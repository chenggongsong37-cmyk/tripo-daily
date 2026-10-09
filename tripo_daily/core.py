"""Core data model, windowing, collection, filtering and report rendering."""
from __future__ import annotations
import hashlib, html, json, os, re, sqlite3
from functools import lru_cache
from difflib import SequenceMatcher
from email.utils import parsedate_to_datetime
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit
from xml.etree import ElementTree as ET
try:
    from zoneinfo import ZoneInfo
except ImportError:  # pragma: no cover
    ZoneInfo = None

CN = ZoneInfo("Asia/Shanghai")
UTC = timezone.utc
COMPETITOR_ALIASES = {
    "Meshy": ("meshy",),
    "影眸科技 / Deemos（平台：Hyper3D；产品：Rodin）": ("deemos", "影眸科技", "影眸", "hyper3d", "hyper 3d", "rodin"),
    "腾讯混元 / Hunyuan3D": ("腾讯混元", "混元3d", "混元 3d", "hunyuan3d", "hunyuan 3d"),
    "阿里 Happy Horse": ("happy horse", "happyhorse"),
}
COMPETITORS = list(COMPETITOR_ALIASES)
DEFAULT_KEYWORDS = ["AI 3D", "3D generation", "三维生成", "3D asset", "世界模型", "world model", "空间智能", "spatial intelligence", "具身智能", "robotics", "simulation", "synthetic data", "算力", "芯片"]
POLICY_KEYWORDS = [
    "出口管制", "出口限制", "实体清单", "贸易限制", "技术封锁", "关税", "制裁",
    "投资审查", "国家安全审查", "政府采购", "数据跨境", "数据出境", "数据本地化",
    "人工智能法", "人工智能监管", "生成式人工智能管理", "算法备案", "模型备案",
    "版权", "著作权", "训练数据", "数据合规", "隐私监管", "反垄断",
    "export control", "export restriction", "entity list", "trade restriction", "tariff",
    "sanction", "investment screening", "national security review", "government procurement",
    "cross-border data", "data localization", "ai regulation", "ai act", "copyright",
    "training data", "privacy regulation", "antitrust", "bis rule", "chip restriction",
]
POLICY_BUSINESS_TERMS = [
    "人工智能", "生成式", "大模型", "模型", "ai ", "3d", "三维", "世界模型", "空间智能",
    "机器人", "仿真", "合成数据", "agent", "gpu", "芯片", "算力", "半导体", "云服务",
    "训练", "推理", "数据中心", "nvidia", "amd", "intel", "tencent", "alibaba", "腾讯", "阿里",
]

def clean_text(value: str | None) -> str:
    value = html.unescape(value or "")
    value = re.sub(r"<[^>]+>", " ", value)
    value = re.sub(r"https?://\S+", "", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value

def clean_editorial_boilerplate(value: str | None) -> str:
    """Remove publisher promos/bylines without altering reported facts."""
    text = clean_text(value)
    patterns = [
        r"#?欢迎关注[^。！？]*?(?:微信公众号|微信号)[^。！？]*[。！？]?",
        r"更多精彩内容[^。！？]*[。！？]?",
        r"(?:文|编译)\s*/\s*[^\s。！？]{1,20}\s*",
        r"来源\s*[：:]\s*[^。！？]{1,80}[。！？]?",
    ]
    for pattern in patterns:
        text = re.sub(pattern, " ", text, flags=re.I)
    return re.sub(r"\s+", " ", text).strip(" #")

def summary_text(item):
    # Full RSS content is stronger evidence than a teaser, which often carries
    # subscription promos. Keep a concise extract rather than dumping content.
    text = clean_editorial_boilerplate(item.content or item.summary)
    if not text or text == item.title or len(text) < 8:
        return f"围绕“{item.title}”的公开信息，当前仅核验到标题和公开摘要，正文细节待进一步核对。"
    sentences = re.split(r"(?<=[。！？.!?])\s+", text)
    selected=[]; total=0
    for sentence in sentences:
        sentence=sentence.strip()
        if not sentence: continue
        selected.append(sentence); total += len(sentence)
        if len(selected) >= 3 or total >= 320: break
    return " ".join(selected)[:420]

_TRANSLATIONS = {
    "artificial intelligence": "人工智能", "generative ai": "生成式人工智能",
    "3d generation": "三维生成", "text-to-3d": "文本生成三维模型",
    "3d model": "三维模型", "3d asset": "三维资产", "world model": "世界模型",
    "spatial intelligence": "空间智能", "robotics": "机器人技术", "simulation": "仿真",
    "synthetic data": "合成数据", "foundation model": "基础模型", "open source": "开源",
    "launch": "发布", "announces": "宣布", "announced": "宣布", "introduces": "推出",
    "new model": "新模型", "research": "研究", "copyright": "版权", "regulation": "监管",
    "export controls": "出口管制", "funding": "融资", "company": "公司", "platform": "平台",
    "available": "可用", "available now": "现已可用", "update": "更新",
}

def builtin_translate(text: str) -> str:
    """Small deterministic translator for no-key deployments.

    It translates common intelligence vocabulary and keeps unknown proper nouns
    intact. This is deliberately conservative: it never invents facts.
    """
    out = clean_text(text)
    for source, target in sorted(_TRANSLATIONS.items(), key=lambda kv: -len(kv[0])):
        out = re.sub(rf"(?i)\b{re.escape(source)}\b", target, out)
    return out

@lru_cache(maxsize=256)
def translate_to_chinese(text: str) -> str:
    """Translate public summary text without requiring an AI model key."""
    text = clean_text(text)[:480]
    if not text or re.search(r"[\u4e00-\u9fff]", text):
        return text
    if os.getenv("TRANSLATION_PROVIDER", "mymemory").lower() == "mymemory":
        try:
            import requests
            response = requests.get(
                "https://api.mymemory.translated.net/get",
                params={"q": text, "langpair": "en|zh-CN"},
                timeout=6,
                headers={"User-Agent": os.getenv("USER_AGENT", "tripo-daily/0.1")},
            )
            response.raise_for_status()
            translated = clean_text(response.json().get("responseData", {}).get("translatedText"))
            if translated and re.search(r"[\u4e00-\u9fff]", translated):
                return translated
        except Exception:
            pass
    return builtin_translate(text)

@lru_cache(maxsize=256)
def translate_to_english(text: str) -> str:
    """Translate a Chinese public summary for the bilingual report."""
    text = clean_text(text)[:480]
    if not text or not re.search(r"[\u4e00-\u9fff]", text):
        return text
    if os.getenv("TRANSLATION_PROVIDER", "mymemory").lower() == "mymemory":
        try:
            import requests
            response = requests.get(
                "https://api.mymemory.translated.net/get",
                params={"q": text, "langpair": "zh-CN|en"},
                timeout=6,
                headers={"User-Agent": os.getenv("USER_AGENT", "tripo-daily/0.1")},
            )
            response.raise_for_status()
            translated = clean_text(response.json().get("responseData", {}).get("translatedText"))
            if translated and len(re.findall(r"[A-Za-z]", translated)) >= 8:
                return translated
        except Exception:
            pass
    return "English translation is temporarily unavailable. Please verify the Chinese summary against the linked public source."

def summaries(item):
    raw = summary_text(item)
    cjk = len(re.findall(r"[\u4e00-\u9fff]", raw))
    latin = len(re.findall(r"[A-Za-z]", raw))
    # A feed may contain a Chinese title or boilerplate while its actual
    # abstract remains English.  Only pass through text when Chinese is the
    # dominant language; otherwise keep the Chinese field explicitly Chinese.
    if cjk >= 12 and cjk >= latin * 0.35:
        return raw, translate_to_english(raw)
    translated = translate_to_chinese(raw)
    if translated != raw and len(re.findall(r"[\u4e00-\u9fff]", translated)) >= 4:
        return f"内置翻译：{translated}", raw
    zh = f"该条目涉及“{builtin_translate(item.title)}”。目前仅获取英文标题或公开摘要；内置翻译未覆盖全部句子，请通过原始链接核验完整内容。"
    return zh, raw

@dataclass
class Item:
    title: str; url: str; source: str; source_type: str = "媒体报道"
    published: datetime | None = None; summary: str = ""; content: str = ""
    competitor: str | None = None; paywall: bool = False; sample: bool = False
    event_time: datetime | None = None; discovered: datetime = field(default_factory=lambda: datetime.now(UTC))
    verified: bool = True; attribution: str = ""; relevance: float = 0.0; reason: str = ""
    impact: str = "待观察"; action: str = "持续跟踪"; confidence: str = "中"
    related_articles: list[dict] = field(default_factory=list)
    @property
    def canonical_url(self):
        p = urlsplit(self.url); return urlunsplit((p.scheme.lower(), p.netloc.lower(), p.path.rstrip("/"), p.query, ""))
    @property
    def id(self): return hashlib.sha256((self.canonical_url + self.title.lower()).encode()).hexdigest()[:24]

def parse_dt(value: str | None) -> datetime | None:
    if not value: return None
    try:
        v = value.strip().replace("Z", "+00:00")
        d = datetime.fromisoformat(v)
        return (d if d.tzinfo else d.replace(tzinfo=UTC)).astimezone(UTC)
    except ValueError:
        try:
            d = parsedate_to_datetime(value)
            return (d if d.tzinfo else d.replace(tzinfo=UTC)).astimezone(UTC)
        except (TypeError, ValueError, OverflowError): return None

def window_for(day: str | datetime, tz=CN):
    if isinstance(day, str): day = datetime.fromisoformat(day).date()
    elif isinstance(day, datetime): day = day.astimezone(tz).date()
    start = datetime(day.year, day.month, day.day, 10, 30, tzinfo=tz).astimezone(UTC)
    return start, start + timedelta(days=1)

def classify(item: Item, keywords=None, competitors=None):
    keywords = keywords or DEFAULT_KEYWORDS; competitors = competitors or COMPETITORS
    text = f"{item.title} {item.summary} {item.content}".lower()
    hits = [k for k in keywords if k.lower() in text]
    policy_hits = [k for k in POLICY_KEYWORDS if k.lower() in text]
    business_hits = [k for k in POLICY_BUSINESS_TERMS if k.lower() in text]
    for c in competitors:
        aliases = COMPETITOR_ALIASES.get(c, (c.lower(),))
        if any(alias.lower() in text for alias in aliases): item.competitor = c
    policy_relevant = bool(policy_hits and business_hits)
    item.relevance = min(1.0, .25 * len(set(hits)) + (.45 if item.competitor else 0) + (.55 if policy_relevant else 0))
    reasons = hits[:4] + ([f"政策影响链：{policy_hits[0]} → {business_hits[0]}"] if policy_relevant else [])
    item.reason = "；".join(reasons) or "未命中配置关键词"
    item.impact = "机会" if any(x in text for x in ("launch", "发布", "突破", "funding", "融资")) else ("风险" if any(x in text for x in ("ban", "control", "监管", "版权", "export")) else "待观察")
    item.confidence = "低" if not item.published or not item.verified else ("高" if item.content else "中")
    return item

def in_window(item, start, end): return item.published is not None and start <= item.published < end

_TOPIC_STOPWORDS = {"the","this","that","with","from","into","built","build","launch","launches","launched","raises","raised","funding","platform","model","models","generation","researcher","company","technology","million","billion","usd","newswire","news","report","reports","announces","announced","unveils","unveiled","assets","agents","artificial","intelligence","intelligent","digital","predictive","cognitive","autonomous","framework","network","ecosystem","ecosystems","accessible","everyone","living"}

def topic_features(item: Item):
    """Extract conservative event features from the headline only."""
    title = clean_text(item.title)
    tokens = re.findall(r"[A-Za-z][A-Za-z0-9-]{3,}", title)
    terms = {w.lower() for w in tokens if w.lower() not in _TOPIC_STOPWORDS}
    # Product/project names usually preserve capitals or contain digits. Do
    # not use arbitrary Chinese fragments or long RSS bodies as entities.
    entities = {w.lower() for w in tokens if (len(w) >= 5 or any(c.isdigit() for c in w)) and (not w.islower() or any(c.isdigit() for c in w)) and w.lower() not in _TOPIC_STOPWORDS}
    competitor_aliases = {a.lower() for aliases in COMPETITOR_ALIASES.values() for a in aliases}
    entities -= competitor_aliases
    concepts = {k.lower() for k in DEFAULT_KEYWORDS + POLICY_KEYWORDS if k.lower() in title.lower()}
    return entities, terms | concepts

def same_topic(a: Item, b: Item) -> bool:
    """Judge event identity without requiring near-identical headlines."""
    an = re.sub(r"\W", "", a.title.lower()); bn = re.sub(r"\W", "", b.title.lower())
    if an and bn and (an in bn or bn in an or SequenceMatcher(None, an, bn).ratio() >= .72): return True
    ae, at = topic_features(a); be, bt = topic_features(b)
    shared_entities = ae & be; shared_terms = at & bt
    close_in_time = bool(a.published and b.published and abs((a.published-b.published).total_seconds()) <= 7*86400)
    # A distinctive shared entity plus time proximity handles different angles
    # on one event (launch, funding, interview). Otherwise require more context.
    return bool(shared_entities and close_in_time and len(shared_terms) >= 1)

def dedupe(items):
    """Merge reports about the same topic while preserving every source link."""
    out=[]; seen_urls={}
    # Newest report becomes the topic's representative; older reports remain
    # attached as evidence and contribute to cumulative volume.
    for i in sorted(items, key=lambda x: (x.published or datetime.min.replace(tzinfo=UTC), x.relevance), reverse=True):
        norm = re.sub(r"\W", "", i.title.lower())
        match = seen_urls.get(i.canonical_url)
        if match is None:
            for candidate in out:
                if same_topic(i, candidate):
                    match = candidate; break
        if match is not None:
            if i.url != match.url and not any(x.get("url") == i.url for x in match.related_articles):
                match.related_articles.append({"title": i.title, "url": i.url, "source": i.source, "source_type": i.source_type, "published": i.published})
            seen_urls[i.canonical_url] = match
            continue
        seen_urls[i.canonical_url] = i; out.append(i)
    return out

def topic_has_window_article(item: Item, start, end) -> bool:
    dates = [item.published] + [x.get("published") for x in item.related_articles]
    return any(d is not None and start <= d < end for d in dates)

def topic_window_volume(item: Item, start, end) -> int:
    dates = [item.published] + [x.get("published") for x in item.related_articles]
    return sum(d is not None and start <= d < end for d in dates)

def topic_stats(item: Item):
    articles = [{"title": item.title, "url": item.url, "source": item.source, "source_type": item.source_type, "published": item.published}] + item.related_articles
    sources = {x.get("source") for x in articles if x.get("source")}
    dates = [x.get("published") for x in articles if x.get("published")]
    span = "时间未完全核验"
    if dates:
        local = sorted(x.astimezone(CN) for x in dates)
        span = local[0].strftime("%m-%d %H:%M") if len(local) == 1 else f"{local[0]:%m-%d %H:%M} ～ {local[-1]:%m-%d %H:%M}"
    return articles, len(sources), span

def topic_digest(item: Item) -> str:
    """One conservative digest for a merged topic, using only fetched text."""
    articles, source_count, _ = topic_stats(item)
    titles = []
    for article in articles:
        title = clean_text(article.get("title"))
        if title and title not in titles: titles.append(title)
    translated = translate_to_chinese(summary_text(item))
    prefix = f"本次采集共发现 {len(articles)} 篇相关文章，来自 {source_count} 个独立来源。"
    return prefix + (f"主要公开信息：{translated}" if translated else f"核心标题：{'；'.join(titles[:3])}")

def report_item_lines(i: Item, heading="###", start=None, end=None):
    """Render one topic consistently in competitor and industry sections."""
    _, en_summary = summaries(i)
    zh_summary = topic_digest(i)
    articles, source_count, span = topic_stats(i)
    def source_link(x):
        when = x["published"].astimezone(CN).strftime("%Y-%m-%d %H:%M") if x.get("published") else "时间未核验"
        return f"[{x['source']}｜{x.get('source_type','来源类型未标注')}｜{when}｜{x['title']}]({x['url']})"
    links = "；".join(source_link(x) for x in articles)
    volume = f"累计相关文章 {len(articles)} 篇"
    if start is not None and end is not None: volume = f"窗口内声量 {topic_window_volume(i,start,end)} 篇｜累计相关文章 {len(articles)} 篇"
    return [f"{heading} {i.title}", f"- 话题数据：{volume}｜独立来源 {source_count} 个｜报道时间 {span}", f"- 来源及来源类型：{i.source}（{i.source_type}）", f"- 原文标题：{i.title}", f"- 相关文章：{links}", f"- 发布时间：{i.published.astimezone(CN).isoformat() if i.published else '时间未核验'}", f"- 核验范围/限制：{'付费墙，仅基于可见摘要；' if i.paywall else ''}{'仅基于公开摘要' if not i.content else '已获取公开摘要'}", f"- 中文报道概述：{zh_summary}", f"- English summary: {en_summary}", f"- 对 Tripo/VAST 的影响（分析判断）：{i.impact}；涉及 {i.reason}。", f"- 建议行动：{i.action}", f"- 置信度：{i.confidence}", ""]

def export_xlsx(items: list[Item], day: str) -> bytes:
    """Export the selected daily topics as a styled, filterable workbook."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook(); ws = wb.active; ws.title = "当日舆情"
    headers = ["媒体名称", "渠道", "区域", "日期", "标题", "链接", "栏目", "主题", "关键内容", "关键词", "竞品", "声量", "独立来源", "影响方向", "影响分析", "建议行动", "置信度", "核验状态"]
    ws.append(headers)
    for i in items:
        articles, source_count, _ = topic_stats(i)
        urls = "\n".join(a["url"] for a in articles)
        topic = i.competitor or ("政策与监管" if "政策影响链" in i.reason else "AI 3D 与行业")
        verified = "仅基于公开摘要" if not i.content else "已获取公开摘要"
        ws.append([i.source, i.source_type, "Global/按来源", i.published.astimezone(CN).date() if i.published else None, i.title, urls, "竞品动态" if i.competitor else "行业动态", topic, topic_digest(i), i.reason, i.competitor or "", len(articles), source_count, i.impact, f"分析判断：{i.impact}；涉及 {i.reason}", i.action, i.confidence, verified])
    dark = PatternFill("solid", fgColor="5B6573"); white = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="8A8A8A")
    for cell in ws[1]: cell.fill=dark; cell.font=white; cell.alignment=Alignment(horizontal="center", vertical="center"); cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(vertical="top", wrap_text=True); cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
        row[0].fill=PatternFill("solid",fgColor="E2E3E5"); row[4].font=Font(bold=True)
    widths=[20,18,14,12,38,42,16,22,55,28,22,10,12,12,35,24,10,18]
    for idx,width in enumerate(widths,1): ws.column_dimensions[get_column_letter(idx)].width=width
    ws.row_dimensions[1].height=28
    for idx in range(2,ws.max_row+1): ws.row_dimensions[idx].height=90
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    if ws.max_row >= 2:
        table=Table(displayName="DailyIntelligence",ref=f"A1:R{ws.max_row}"); table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showFirstColumn=False,showLastColumn=False); ws.add_table(table)
    ws.sheet_view.showGridLines=False; ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1
    info=wb.create_sheet("说明"); info.append(["Tripo AI 舆情与行业日报", day]); info.append(["时效口径", "北京时间前一天 10:30 至当天 10:30"]); info.append(["声量口径", "本次采集范围内合并到同一话题的公开文章数量，不代表全网总量"]); info.column_dimensions["A"].width=22; info.column_dimensions["B"].width=80
    stream=BytesIO(); wb.save(stream); return stream.getvalue()

class Store:
    def __init__(self, path):
        Path(path).parent.mkdir(parents=True, exist_ok=True); self.db=sqlite3.connect(path)
        self.db.execute("create table if not exists items(id text primary key, data text not null)"); self.db.commit()
    def put(self, item): self.db.execute("insert or ignore into items values (?,?)", (item.id, json.dumps(item.__dict__, default=lambda x:x.isoformat() if isinstance(x,datetime) else x, ensure_ascii=False))); self.db.commit()

def parse_feed(xml: str | bytes, source: str, source_type="媒体报道"):
    root=ET.fromstring(xml); items=[]
    for e in root.findall(".//item") + root.findall(".//{http://www.w3.org/2005/Atom}entry"):
        def val(*names):
            for n in names:
                x = e.find(n)
                if x is None:
                    x = e.find("{http://www.w3.org/2005/Atom}" + n)
                if x is None:
                    x = next((node for node in e if node.tag.rsplit("}", 1)[-1] == n), None)
                if x is not None: return (x.text or "").strip()
            return ""
        link=val("link") or next((x.attrib.get("href","") for x in e.findall("*") if x.tag.endswith("link")), "")
        summary = val("description","summary")
        content = val("encoded","content")
        items.append(Item(val("title") or "(无标题)", link, source, source_type, parse_dt(val("pubDate","published","updated")), summary, content))
    return items

def render_report(day, start, end, items, status, failures=(), pending=()):
    d=start.astimezone(CN); weekday="一二三四五六日"[d.weekday()]
    lines=[f"# {d.year}年{d.month}月{d.day}日 星期{weekday}｜Tripo AI 舆情与行业日报", "", f"覆盖窗口：{start.astimezone(CN):%Y-%m-%d %H:%M} ～ {end.astimezone(CN):%Y-%m-%d %H:%M}（北京时间）", f"生成时间：{datetime.now(UTC).astimezone(CN):%Y-%m-%d %H:%M}", f"采集状态：成功来源数 {status.get('ok',0)}，失败来源数 {status.get('failed',0)}，候选数 {status.get('candidates',0)}，入选数 {len(items)}" + ("；覆盖不完整" if failures else ""), "", "## 今日要点"]
    for i in items[:5]: lines.append(f"- [{i.title}]({i.url}) — {i.impact}：{i.reason}")
    lines += ["", "## 竞品动态追踪"]
    for c in COMPETITORS:
        xs=[i for i in items if i.competitor==c]
        pending_xs = [i for i in pending if i.competitor == c]
        lines.append(f"### {c}")
        if xs:
            for i in xs: lines.extend(report_item_lines(i, "####", start, end))
        else:
            lines.append("本窗口未发现可核验的新动态（不等于确定没有动态）。")
        if pending_xs:
            for i in pending_xs[:5]:
                articles, source_count, _ = topic_stats(i)
                links = "；".join(f"[{x['source']}]({x['url']})" for x in articles)
                lines.append(f"- 待核验话题：{i.title}｜声量 {len(articles)} 篇｜独立来源 {source_count} 个｜{links}")
    lines += ["", "## AI 3D 与行业动态"]
    industry_items = [i for i in items if not i.competitor]
    if industry_items:
        for i in industry_items: lines.extend(report_item_lines(i, "###", start, end))
    else:
        lines.append("本窗口未发现竞品栏目之外的可核验行业动态。")
    lines += ["## 国际与国家级重大事件", "仅保留能解释产品、市场、合规、供应链或算力影响链条的窗口内条目；窗口外旧消息不展示。", "", "## 待核验线索", "仅保留发布时间未能核验的线索；已确认属于窗口外的旧消息会被排除，不进入当前日报。"]
    groups = {"竞品与产品": [], "AI 3D、世界模型与空间智能": [], "政策、版权与监管": [], "算力、芯片与基础设施": [], "其他行业": []}
    for x in pending:
        text = f"{x.title} {x.summary}".lower()
        if x.competitor: continue
        if any(k in text for k in ("3d", "三维", "world model", "世界模型", "空间智能", "spatial")): key = "AI 3D、世界模型与空间智能"
        elif any(k in text for k in ("regulation", "版权", "copyright", "监管", "ai act", "export control", "政策")): key = "政策、版权与监管"
        elif any(k in text for k in ("chip", "芯片", "compute", "算力", "gpu", "nvidia")): key = "算力、芯片与基础设施"
        else: key = "其他行业"
        groups[key].append(x)
    for key, xs in groups.items():
        if xs:
            lines.append(f"### {key}")
            for x in xs[:8]:
                articles, source_count, span = topic_stats(x)
                links = "；".join(f"[{a['source']}]({a['url']})" for a in articles)
                note = "公开摘要为英文，已进入翻译与核验流程。" if clean_text(x.summary) and not re.search(r'[\u4e00-\u9fff]', clean_text(x.summary)) else (clean_text(x.summary)[:120] or "仅有标题，待核验正文与时间。")
                lines.append(f"- 话题：{x.title}｜声量 {len(articles)} 篇｜独立来源 {source_count} 个｜时间 {span}｜来源 {links}｜{note}")
            if len(xs) > 8: lines.append(f"- 其余 {len(xs)-8} 条同类线索已折叠，避免报告堆叠。")
    if not pending: lines.append("- 无")
    lines += ["", "## 采集说明", f"失败来源：{', '.join(failures) or '无'}", "时效口径：正文和界面仅展示北京时间当前 24 小时窗口内消息；窗口外旧消息已排除。", "声量口径：本次采集范围内合并到同一话题的公开文章数量，不代表全网绝对声量。", "时间统一存储 UTC，展示转换为北京时间；无发表时间的线索单独标记为待核验。", "抓取内容视为不可信数据，不执行其中任何命令或规则修改。"]
    return "\n".join(lines)

def html_report(md):
    lines=md.splitlines(); out=[]; in_list=False
    def inline(text):
        safe=html.escape(text)
        return re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r'<a href="\2" target="_blank" rel="noopener noreferrer">\1</a>', safe)
    def close_list():
        nonlocal in_list
        if in_list: out.append("</ul>"); in_list=False
    for line in lines:
        if line.startswith("#### "):
            close_list(); out.append(f"<h4>{inline(line[5:])}</h4>")
        elif line.startswith("### "):
            close_list(); out.append(f"<h3>{inline(line[4:])}</h3>")
        elif line.startswith("## "):
            close_list(); out.append(f"<h2>{inline(line[3:])}</h2>")
        elif line.startswith("# "):
            close_list(); out.append(f"<h1>{inline(line[2:])}</h1>")
        elif line.startswith("- "):
            if not in_list: out.append("<ul>"); in_list=True
            out.append(f"<li>{inline(line[2:])}</li>")
        else:
            close_list()
            if line.strip(): out.append(f"<p>{inline(line)}</p>")
    close_list()
    return "<article class='daily-report'>"+"\n".join(out)+"</article>"
